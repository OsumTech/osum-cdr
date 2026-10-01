from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.billing import bill_call, bill_inbound
from app.csv_reconciliation import compare_csv
from app.models import Call, CdrPartition, Did, Integration, Ledger, SipAccount, Tenant, User, db, utcnow
from app.reconciliation import approve, canonical_rows, fetch_partition, save_snapshot

DAY = date(2025, 1, 1)


@pytest.fixture
def inbound_setup(app):
    did = Did(tenant_id=1, number='+442071234567', monthly_charge=5, next_billing_date=date(2027, 1, 1), billing_day=1)
    other = Did(tenant_id=2, number='+442071234568', monthly_charge=5, next_billing_date=date(2027, 1, 1), billing_day=1)
    owner = User(email='inbound-owner@example.test', is_admin=True, password_hash='unused')
    db.session.add_all([did, other, owner, Integration(id=1, reconciliation_enabled=True, currency_confirmed=True, token_version=1)])
    db.session.commit()
    return did, other, owner.id


def row(**changes):
    return {'type': 'incoming', 'did_number': '442071234567', 'to': 'extension-100', 'from': 'anonymous',
            'timestamp': '2025-01-01T12:00:00Z', 'duration': 30, 'amount': '0.0900', **changes}


def ready(did, rows):
    groups = canonical_rows(rows, did, DAY)
    save_snapshot(did, DAY, groups, 1)
    part = db.session.scalar(select(CdrPartition).where(CdrPartition.did_id == did.id))
    part.checked_at = utcnow() - timedelta(minutes=6)
    db.session.commit()
    save_snapshot(did, DAY, groups, 1)
    return part


def test_inbound_exact_cost_mapping_idempotency_and_outbound_unchanged(inbound_setup):
    did, other, owner = inbound_setup
    tenant = db.session.get(Tenant, 1)
    tenant.mobile_rate, tenant.billing_increment = Decimal('9'), 60
    db.session.commit()
    mixed = [row(), row(), row(did_number=other.number)]
    part = ready(did, mixed)
    assert part.record_count == 2 and part.wholesale_total == Decimal('0.18')
    assert approve(part.id, part.digest, owner) == 2
    assert approve(part.id, part.digest, owner) == 0
    assert tenant.balance == Decimal('-0.18')
    assert db.session.get(Tenant, 2).balance == 0
    incoming = db.session.scalars(select(Call)).all()
    assert all(c.did_id == did.id and c.sip_account_id is None and c.cost == c.wholesale_cost == Decimal('0.09') for c in incoming)
    before = tenant.balance
    assert bill_call(1, 'outbound-after-inbound', datetime(2025, 1, 1, 12, tzinfo=timezone.utc), '447700900123', 30)
    assert tenant.balance == before - Decimal('9')
    updated = ready(did, [row(amount='0.1'), row(amount='0.1')])
    assert updated.status == 'correction'
    with pytest.raises(ValueError):
        approve(updated.id, updated.digest, owner)


def test_missing_did_or_cost_fails_and_zero_duration_charge_is_kept(inbound_setup):
    did, _, owner = inbound_setup
    with pytest.raises(ValueError):
        canonical_rows([row(did_number=None)], did, DAY)
    with pytest.raises(ValueError, match='charge'):
        canonical_rows([row(amount=None)], did, DAY)
    part = ready(did, [row(duration=0), row(duration=0, amount=0)])
    assert part.record_count == 1
    approve(part.id, part.digest, owner)
    assert db.session.get(Tenant, 1).balance == Decimal('-0.09')
    assert db.session.scalar(select(Call)).duration == 0


def test_inbound_complete_fetch_has_no_sip_filter_and_reuses_day(inbound_setup, monkeypatch):
    did, other, _ = inbound_setup
    requests = []
    def fake(token, params):
        requests.append(params)
        assert params['type'] == 'incoming' and params['missed'] == '1'
        assert 'sip_account' not in params
        return {'calls': [row(), row(did_number=other.number)], 'pagination': {'page': 1, 'per_page': 1000, 'total_pages': 1, 'total_records': 2}}
    monkeypatch.setattr('app.reconciliation.read_calls', fake)
    cache = {}
    assert len(fetch_partition('unused', did, DAY, cache)) == 1
    assert len(fetch_partition('unused', other, DAY, cache)) == 1
    assert len(requests) == 1


def test_inbound_csv_and_excel_tenant_isolation(inbound_setup, signed_in):
    did, other, owner = inbound_setup
    part = ready(did, [row()])
    csv = b'Date,Time,SIP ID,Type,From,To,Duration,Charge\n01/01/25,12:00:00 pm,,INCOMING,anonymous,442071234567,30,0.09\n'
    assert compare_csv(csv, part, 0)['matched']
    approve(part.id, part.digest, owner)
    bill_inbound(other.id, 'other-inbound', datetime(2025, 1, 1, 12, tzinfo=timezone.utc), 30, '0.25')
    response = signed_in.get('/calls/export.xlsx?start=2025-01-01&end=2025-01-01')
    sheet = load_workbook(BytesIO(response.data)).active
    assert sheet.max_row == 6
    assert sheet['B6'].value == did.number
    assert sheet['C6'].value == did.number
    assert sheet['F6'].value == 0.09
    assert b'Inbound' in signed_in.get('/calls?start=2025-01-01&end=2025-01-01').data
